"""Dossier de compétence consultant — export PDF & DOCX (EPIC J)."""
from io import BytesIO

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

_KAIROS = colors.HexColor("#D9532F")
_INK = colors.HexColor("#1A1710")
_MUTE = colors.HexColor("#8A8272")


def _yr(e: dict, key: str = "year") -> str:
    return f" ({e[key]})" if e.get(key) else ""


def build_dossier_pdf(p) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm, topMargin=1.4 * cm, bottomMargin=1.4 * cm)
    s = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=s["Title"], fontSize=20, textColor=_INK, alignment=0, spaceAfter=2)
    role = ParagraphStyle("role", parent=s["Normal"], fontSize=11, textColor=_KAIROS, spaceAfter=2)
    meta = ParagraphStyle("meta", parent=s["Normal"], fontSize=9, textColor=_MUTE)
    sec = ParagraphStyle("sec", parent=s["Heading2"], fontSize=12, textColor=_KAIROS, spaceBefore=12, spaceAfter=4)
    body = ParagraphStyle("body", parent=s["Normal"], fontSize=9.5, leading=13)

    e = [Paragraph(p.name, h1)]
    if p.headline or p.role:
        e.append(Paragraph(p.headline or p.role, role))
    m = " &nbsp;·&nbsp; ".join(x for x in [p.seniority, p.location, p.email] if x)
    if m:
        e.append(Paragraph(m, meta))
    if p.summary:
        e += [Paragraph("Résumé", sec), Paragraph(p.summary, body)]
    if p.skills:
        e += [Paragraph("Compétences", sec), Paragraph(" &nbsp;·&nbsp; ".join(p.skills), body)]
    if p.experiences:
        e.append(Paragraph("Expériences", sec))
        for x in p.experiences:
            t = f"<b>{x.get('title', '')}</b>"
            if x.get("company"):
                t += f" — {x['company']}"
            if x.get("period"):
                t += f" <font color='#8A8272'>({x['period']})</font>"
            e.append(Paragraph(t, body))
            if x.get("description"):
                e.append(Paragraph(x["description"], body))
            e.append(Spacer(1, 4))
    if p.education:
        e.append(Paragraph("Formations", sec))
        for x in p.education:
            e.append(Paragraph(f"<b>{x.get('degree', '')}</b> — {x.get('school', '')}{_yr(x)}", body))
    if p.certifications:
        e.append(Paragraph("Certifications", sec))
        for x in p.certifications:
            e.append(Paragraph(f"{x.get('name', '')} — {x.get('issuer', '')}{_yr(x)}", body))
    if p.languages:
        e.append(Paragraph("Langues", sec))
        e.append(Paragraph(" &nbsp;·&nbsp; ".join(f"{x.get('name', '')} ({x.get('level', '')})" for x in p.languages), body))
    doc.build(e)
    return buf.getvalue()


def build_dossier_docx(p) -> bytes:
    from docx import Document
    from docx.shared import RGBColor

    d = Document()
    d.add_heading(p.name, level=0)
    if p.headline or p.role:
        run = d.add_paragraph().add_run(p.headline or p.role)
        run.bold = True
        run.font.color.rgb = RGBColor(0xD9, 0x53, 0x2F)
    m = " · ".join(x for x in [p.seniority, p.location, p.email] if x)
    if m:
        d.add_paragraph(m)

    def section(title):
        d.add_heading(title, level=1)

    if p.summary:
        section("Résumé")
        d.add_paragraph(p.summary)
    if p.skills:
        section("Compétences")
        d.add_paragraph(" · ".join(p.skills))
    if p.experiences:
        section("Expériences")
        for x in p.experiences:
            para = d.add_paragraph()
            r = para.add_run(f"{x.get('title', '')} — {x.get('company', '')}")
            r.bold = True
            if x.get("period"):
                para.add_run(f" ({x['period']})")
            if x.get("description"):
                d.add_paragraph(x["description"])
    if p.education:
        section("Formations")
        for x in p.education:
            d.add_paragraph(f"{x.get('degree', '')} — {x.get('school', '')}{_yr(x)}")
    if p.certifications:
        section("Certifications")
        for x in p.certifications:
            d.add_paragraph(f"{x.get('name', '')} — {x.get('issuer', '')}{_yr(x)}")
    if p.languages:
        section("Langues")
        d.add_paragraph(" · ".join(f"{x.get('name', '')} ({x.get('level', '')})" for x in p.languages))
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()
